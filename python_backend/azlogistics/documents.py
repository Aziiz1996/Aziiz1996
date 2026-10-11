from __future__ import annotations
from io import BytesIO
from urllib.parse import urljoin, urlparse, quote
import hashlib
import ipaddress
import socket
import uuid
import httpx
from pypdf import PdfReader
from pptx import Presentation
from .config import settings
from .firebase import bucket

ALLOWED_MIME = {"application/pdf": ".pdf", "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx"}

def assert_public_https(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme != "https" or not parsed.hostname or parsed.username or parsed.password:
        raise ValueError("unsafe-url")
    for result in socket.getaddrinfo(parsed.hostname, 443, type=socket.SOCK_STREAM):
        ip = ipaddress.ip_address(result[4][0])
        if not ip.is_global:
            raise ValueError("private-address")

async def download(url: str) -> tuple[bytes, str]:
    max_bytes = settings().max_file_mb * 1024 * 1024
    current = url
    async with httpx.AsyncClient(timeout=45, follow_redirects=False, headers={"User-Agent": "AzLogisticsAcademicBot/1.0"}) as client:
        for _ in range(5):
            assert_public_https(current)
            async with client.stream("GET", current) as response:
                if response.status_code in {301, 302, 303, 307, 308}:
                    current = urljoin(current, response.headers["location"])
                    continue
                response.raise_for_status()
                declared = response.headers.get("content-type", "").split(";")[0].lower()
                parts, size = [], 0
                async for chunk in response.aiter_bytes():
                    size += len(chunk)
                    if size > max_bytes:
                        raise ValueError("file-too-large")
                    parts.append(chunk)
                content = b"".join(parts)
                mime = "application/pdf" if content.startswith(b"%PDF-") else "application/vnd.openxmlformats-officedocument.presentationml.presentation" if content.startswith(b"PK\x03\x04") and (declared.endswith("presentationml.presentation") or urlparse(current).path.lower().endswith(".pptx")) else ""
                if mime not in ALLOWED_MIME:
                    raise ValueError("unsupported-file")
                return content, mime
    raise ValueError("too-many-redirects")

def extract_text(content: bytes, mime: str) -> str:
    if mime == "application/pdf":
        pages = PdfReader(BytesIO(content)).pages[:300]
        text = "\n".join(page.extract_text() or "" for page in pages)
    else:
        presentation = Presentation(BytesIO(content))
        text = "\n".join(shape.text for slide in presentation.slides[:300] for shape in slide.shapes if hasattr(shape, "text"))
    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if len(text) < 800:
        raise ValueError("insufficient-text")
    return text[:500000]

def sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()

def upload_document(course_id: str, digest: str, content: bytes, mime: str, text: str) -> dict:
    storage = bucket()
    extension = ALLOWED_MIME[mime]
    token = str(uuid.uuid4())
    path = f"courses/{course_id}/references/{digest}{extension}"
    blob = storage.blob(path)
    blob.metadata = {"firebaseStorageDownloadTokens": token}
    blob.upload_from_string(content, content_type=mime)
    text_path = f"courses/{course_id}/extracted/{digest}.txt"
    storage.blob(text_path).upload_from_string(text, content_type="text/plain; charset=utf-8")
    public_url = f"https://firebasestorage.googleapis.com/v0/b/{storage.name}/o/{quote(path, safe='')}?alt=media&token={token}"
    return {"storagePath": path, "textStoragePath": text_path, "downloadUrl": public_url}

def load_extracted_text(path: str) -> str:
    return bucket().blob(path).download_as_text(encoding="utf-8")
